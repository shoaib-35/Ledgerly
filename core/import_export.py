"""Safe transaction import/export for CSV and Excel workbooks.

The import format is deliberately explicit. Exported files contain both stable
UUIDs and human-readable names so a Ledgerly export can be re-imported without
relying on names alone. New files can be created from the downloadable template
using names instead of UUIDs; duplicate names are rejected unless an ID is
provided.
"""

import csv
import io
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from uuid import UUID

from django.db import transaction as db_transaction
from django.http import HttpResponse
from django.utils import timezone

from .models import Account, Person, Transaction
from .services import FinancialValidationError, create_transaction, void_transaction

try:
    import openpyxl
except ImportError:  # pragma: no cover - dependency is declared in requirements.txt
    openpyxl = None

try:
    import xlrd
except ImportError:  # pragma: no cover
    xlrd = None

try:
    import xlwt
except ImportError:  # pragma: no cover
    xlwt = None


FORMAT_COLUMNS = (
    'transaction_id',
    'date',
    'transaction_type',
    'amount',
    'description',
    'status',
    'person_id',
    'person_name',
    'account_id',
    'account_name',
    'source_account_id',
    'source_account_name',
    'destination_account_id',
    'destination_account_name',
    'payment_source',
    'created_at',
)

REQUIRED_COLUMNS = {'date', 'transaction_type', 'amount', 'description'}
MAX_IMPORT_BYTES = 10 * 1024 * 1024
MAX_IMPORT_ROWS = 10_000


def _safe_text(value):
    """Prevent spreadsheet formula injection while preserving normal text."""
    if value is None:
        return ''
    text = str(value)
    if text.startswith(('=', '+', '-', '@')):
        return "'" + text
    return text


def _restore_text(value):
    if value is None:
        return ''
    text = str(value).strip()
    if len(text) > 1 and text[0] == "'" and text[1] in '=+-@':
        return text[1:]
    return text


def _normal_header(value):
    return re.sub(r'\s+', '_', str(value or '').strip().lower())


def _read_csv(data):
    text = None
    for encoding in ('utf-8-sig', 'utf-8', 'utf-16'):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise FinancialValidationError('The CSV file must be UTF-8 or UTF-16 encoded.')
    reader = csv.reader(io.StringIO(text))
    rows = list(reader)
    if not rows:
        raise FinancialValidationError('The import file is empty.')
    headers = [_normal_header(value) for value in rows[0]]
    return headers, [dict(zip(headers, row)) for row in rows[1:] if any(str(v).strip() for v in row)]


def _read_xlsx(data):
    if openpyxl is None:
        raise FinancialValidationError('XLSX support is unavailable. Install openpyxl.')
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise FinancialValidationError('The XLSX file could not be opened.') from exc
    sheet = workbook.active
    values = list(sheet.iter_rows(values_only=True))
    workbook.close()
    if not values:
        raise FinancialValidationError('The import file is empty.')
    headers = [_normal_header(value) for value in values[0]]
    return headers, [dict(zip(headers, row)) for row in values[1:] if any(v not in (None, '') for v in row)]


def _read_xls(data):
    if xlrd is None:
        raise FinancialValidationError('XLS support is unavailable. Install xlrd.')
    try:
        workbook = xlrd.open_workbook(file_contents=data, on_demand=True)
        sheet = workbook.sheet_by_index(0)
        if sheet.nrows == 0:
            raise FinancialValidationError('The import file is empty.')
        headers = [_normal_header(sheet.cell_value(0, col)) for col in range(sheet.ncols)]
        rows = []
        for row_idx in range(1, sheet.nrows):
            values = [sheet.cell_value(row_idx, col) for col in range(sheet.ncols)]
            if any(value not in ('', None) for value in values):
                rows.append(dict(zip(headers, values)))
        workbook.release_resources()
        return headers, rows
    except FinancialValidationError:
        raise
    except Exception as exc:
        raise FinancialValidationError('The XLS file could not be opened.') from exc


def read_import_file(uploaded_file):
    if uploaded_file.size > MAX_IMPORT_BYTES:
        raise FinancialValidationError('Import files must be 10 MB or smaller.')
    name = uploaded_file.name.lower()
    data = uploaded_file.read()
    if name.endswith('.csv'):
        return _read_csv(data)
    if name.endswith('.xlsx'):
        return _read_xlsx(data)
    if name.endswith('.xls'):
        return _read_xls(data)
    raise FinancialValidationError('Unsupported file type. Use .xlsx, .xls, or .csv.')


def _validate_headers(headers):
    if len(headers) != len(set(headers)):
        raise FinancialValidationError('The import file contains duplicate column headers.')
    missing = REQUIRED_COLUMNS - set(headers)
    if missing:
        raise FinancialValidationError('Missing required columns: ' + ', '.join(sorted(missing)) + '.')
    unknown = set(headers) - set(FORMAT_COLUMNS)
    if unknown:
        raise FinancialValidationError('Unknown columns: ' + ', '.join(sorted(unknown)) + '.')


def _cell(row, key):
    return _restore_text(row.get(key, ''))


def _parse_date(value, row_number):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        # Excel serial dates are only accepted for spreadsheet formats. This
        # branch is intentionally conservative; normal exports use ISO text.
        try:
            return date(1899, 12, 30) + __import__('datetime').timedelta(days=float(value))
        except (OverflowError, ValueError):
            pass
    text = _cell({'value': value}, 'value') if not isinstance(value, str) else value.strip()
    for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y', '%Y/%m/%d'):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise FinancialValidationError(f'Row {row_number}: date must use YYYY-MM-DD.')


def _parse_datetime(value, row_number):
    if not value:
        return None
    if isinstance(value, datetime):
        return timezone.make_aware(value) if timezone.is_naive(value) else value
    text = str(value).strip().replace('Z', '+00:00')
    try:
        parsed = datetime.fromisoformat(text)
    except ValueError as exc:
        raise FinancialValidationError(f'Row {row_number}: created_at must be an ISO datetime.') from exc
    return timezone.make_aware(parsed) if timezone.is_naive(parsed) else parsed


def _parse_amount(value, row_number):
    text = str(value).strip().replace(',', '')
    try:
        amount = Decimal(text).quantize(Decimal('0.01'))
    except (InvalidOperation, ValueError, TypeError) as exc:
        raise FinancialValidationError(f'Row {row_number}: amount must be a valid number.') from exc
    if amount <= 0:
        raise FinancialValidationError(f'Row {row_number}: amount must be greater than zero.')
    return amount


def _normalize_choice(value, choices, field_name, row_number, default=None):
    text = str(value or '').strip()
    if not text and default is not None:
        return default
    for choice_value, label in choices:
        if text.upper() == choice_value.upper() or text.casefold() == str(label).casefold():
            return choice_value
    allowed = ', '.join(value for value, _ in choices)
    raise FinancialValidationError(f'Row {row_number}: {field_name} must be one of {allowed}.')


def _resolve_by_id_or_name(model, user, row, id_key, name_key, label, row_number, active_only=True):
    raw_id = _cell(row, id_key)
    raw_name = _cell(row, name_key)
    qs = model.objects.for_user(user)
    if active_only and hasattr(model, 'Status'):
        qs = qs.filter(status=model.Status.ACTIVE)
    if raw_id:
        try:
            obj = qs.get(pk=UUID(raw_id))
        except (ValueError, model.DoesNotExist) as exc:
            raise FinancialValidationError(f'Row {row_number}: {label} ID is not valid for this user.') from exc
        if raw_name and obj.name != raw_name:
            raise FinancialValidationError(f'Row {row_number}: {label} ID and name do not match.')
        return obj
    if not raw_name:
        return None
    matches = list(qs.filter(name__iexact=raw_name)[:2])
    if not matches:
        raise FinancialValidationError(f'Row {row_number}: {label} "{raw_name}" was not found.')
    if len(matches) > 1:
        raise FinancialValidationError(
            f'Row {row_number}: more than one {label.lower()} is named "{raw_name}". Use its ID.'
        )
    return matches[0]


def _normalize_row(row, user, row_number):
    transaction_id = _cell(row, 'transaction_id')
    existing = None
    if transaction_id:
        try:
            existing = Transaction.objects.for_user(user).filter(pk=UUID(transaction_id)).first()
        except ValueError as exc:
            raise FinancialValidationError(f'Row {row_number}: transaction_id is not a valid UUID.') from exc
        if existing:
            return {'duplicate': True, 'transaction': existing}

    transaction_type = _normalize_choice(
        row.get('transaction_type'), Transaction.TransactionType.choices,
        'transaction_type', row_number
    )
    status = _normalize_choice(
        row.get('status'), Transaction.Status.choices, 'status', row_number,
        default=Transaction.Status.ACTIVE,
    )
    amount = _parse_amount(row.get('amount'), row_number)
    description = _cell(row, 'description')
    if not description:
        raise FinancialValidationError(f'Row {row_number}: description is required.')

    person = _resolve_by_id_or_name(Person, user, row, 'person_id', 'person_name', 'Person', row_number)
    account = _resolve_by_id_or_name(Account, user, row, 'account_id', 'account_name', 'Account', row_number)
    source = _resolve_by_id_or_name(
        Account, user, row, 'source_account_id', 'source_account_name', 'Source account', row_number
    )
    destination = _resolve_by_id_or_name(
        Account, user, row, 'destination_account_id', 'destination_account_name', 'Destination account', row_number
    )

    if transaction_type in (Transaction.TransactionType.EXPENSE, Transaction.TransactionType.RECEIVED) and person is None:
        raise FinancialValidationError(f'Row {row_number}: person is required for {transaction_type}.')
    if account is None and transaction_type != Transaction.TransactionType.TRANSFER:
        raise FinancialValidationError(f'Row {row_number}: account is required.')

    payment_source = _normalize_choice(
        row.get('payment_source'),
        (('BANK_ACCOUNT', 'Bank account'), ('CASH', 'Cash')),
        'payment_source', row_number,
        default=None,
    )

    if transaction_type == Transaction.TransactionType.TRANSFER:
        if account is not None and destination is not None and account.pk != destination.pk:
            raise FinancialValidationError(f'Row {row_number}: account and destination_account do not match.')
        destination = destination or account
        if source is None or destination is None:
            raise FinancialValidationError(f'Row {row_number}: transfer requires source and destination accounts.')
        if source.pk == destination.pk:
            raise FinancialValidationError(f'Row {row_number}: transfer source and destination must differ.')
        account = destination
    elif transaction_type == Transaction.TransactionType.CARD_PAYMENT:
        if account is None:
            raise FinancialValidationError(f'Row {row_number}: card payment requires a credit-card account.')
        if payment_source == 'BANK_ACCOUNT' and source is None:
            raise FinancialValidationError(f'Row {row_number}: bank-account card payments require source_account.')
        if payment_source == 'CASH' and source is not None:
            raise FinancialValidationError(f'Row {row_number}: cash card payments cannot have source_account.')
        if payment_source is None:
            raise FinancialValidationError(f'Row {row_number}: card payment requires payment_source.')

    created_at = _parse_datetime(row.get('created_at'), row_number)
    return {
        'duplicate': False,
        'date': _parse_date(row.get('date'), row_number),
        'transaction_type': transaction_type,
        'amount': amount,
        'description': description,
        'status': status,
        'person_id': person.pk if person else None,
        'account_id': account.pk if account else None,
        'payment_source': payment_source,
        'source_account_id': source.pk if source else None,
        'destination_account_id': destination.pk if destination else None,
        'created_at': created_at,
    }


def import_transactions(uploaded_file, user):
    headers, rows = read_import_file(uploaded_file)
    _validate_headers(headers)
    if len(rows) > MAX_IMPORT_ROWS:
        raise FinancialValidationError(f'Import files may contain at most {MAX_IMPORT_ROWS} transactions.')

    normalized = []
    duplicates = 0
    for index, row in enumerate(rows, start=2):
        try:
            parsed = _normalize_row(row, user, index)
        except FinancialValidationError:
            raise
        if parsed['duplicate']:
            duplicates += 1
        else:
            normalized.append(parsed)

    # Import in file order. For equal dates this preserves the exported
    # created_at order and therefore the same-day balance sequence.
    imported = []
    with db_transaction.atomic():
        for item in normalized:
            transaction = create_transaction(
                user=user,
                transaction_type=item['transaction_type'],
                transaction_date=item['date'],
                person_id=item['person_id'],
                account_id=item['account_id'],
                amount=item['amount'],
                description=item['description'],
                payment_source=item['payment_source'],
                source_account_id=item['source_account_id'],
            )
            if item['created_at'] is not None:
                transaction.created_at = item['created_at']
                transaction.save(update_fields=('created_at',))
            if item['status'] == Transaction.Status.VOIDED:
                void_transaction(user=user, transaction_id=transaction.pk)
            imported.append(transaction)
    return len(imported), duplicates


def _export_rows(transactions):
    rows = []
    for transaction in transactions:
        details = None
        payment = None
        if transaction.transaction_type == Transaction.TransactionType.TRANSFER:
            try:
                details = transaction.money_transfer_details
            except Exception:
                details = None
        if transaction.transaction_type == Transaction.TransactionType.CARD_PAYMENT:
            try:
                payment = transaction.card_payment_details
            except Exception:
                payment = None
        destination = details.destination_account if details else None
        source = details.source_account if details else (payment.source_account if payment else None)
        rows.append({
            'transaction_id': str(transaction.pk),
            'date': transaction.date.isoformat(),
            'transaction_type': transaction.transaction_type,
            'amount': str(transaction.amount),
            'description': _safe_text(transaction.description),
            'status': transaction.status,
            'person_id': str(transaction.person_id) if transaction.person_id else '',
            'person_name': _safe_text(transaction.person.name) if transaction.person else '',
            'account_id': str(transaction.account_id) if transaction.account_id else '',
            'account_name': _safe_text(transaction.account.name) if transaction.account else '',
            'source_account_id': str(source.pk) if source else '',
            'source_account_name': _safe_text(source.name) if source else '',
            'destination_account_id': str(destination.pk) if destination else '',
            'destination_account_name': _safe_text(destination.name) if destination else '',
            'payment_source': payment.payment_source if payment else '',
            'created_at': transaction.created_at.isoformat(),
        })
    return rows


def export_csv(transactions):
    response = HttpResponse(content_type='text/csv; charset=utf-8')
    response['Content-Disposition'] = 'attachment; filename="ledgerly_transactions.csv"'
    response.write('\ufeff')
    writer = csv.DictWriter(response, fieldnames=FORMAT_COLUMNS)
    writer.writeheader()
    writer.writerows(_export_rows(transactions))
    return response


def export_xlsx(transactions):
    if openpyxl is None:
        raise FinancialValidationError('XLSX support is unavailable. Install openpyxl.')
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = 'Transactions'
    sheet.append(list(FORMAT_COLUMNS))
    for row in _export_rows(transactions):
        sheet.append([row[column] for column in FORMAT_COLUMNS])
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = openpyxl.styles.Font(bold=True)
    for column in sheet.columns:
        letter = column[0].column_letter
        max_len = min(max(len(str(cell.value or '')) for cell in column) + 2, 42)
        sheet.column_dimensions[letter].width = max_len
    output = io.BytesIO()
    workbook.save(output)
    response = HttpResponse(
        output.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = 'attachment; filename="ledgerly_transactions.xlsx"'
    return response


def export_xls(transactions):
    if xlwt is None:
        raise FinancialValidationError('XLS support is unavailable. Install xlwt.')
    workbook = xlwt.Workbook()
    sheet = workbook.add_sheet('Transactions')
    header_style = xlwt.easyxf('font: bold on; pattern: pattern solid, fore_colour gray25;')
    for col, name in enumerate(FORMAT_COLUMNS):
        sheet.write(0, col, name, header_style)
    for row_index, row in enumerate(_export_rows(transactions), start=1):
        for col, name in enumerate(FORMAT_COLUMNS):
            sheet.write(row_index, col, row[name])
    for col, name in enumerate(FORMAT_COLUMNS):
        sheet.col(col).width = min(max(len(name) + 2, 12), 36) * 256
    output = io.BytesIO()
    workbook.save(output)
    response = HttpResponse(output.getvalue(), content_type='application/vnd.ms-excel')
    response['Content-Disposition'] = 'attachment; filename="ledgerly_transactions.xls"'
    return response


def export_transactions(transactions, file_format):
    file_format = (file_format or '').lower()
    if file_format == 'csv':
        return export_csv(transactions)
    if file_format == 'xlsx':
        return export_xlsx(transactions)
    if file_format == 'xls':
        return export_xls(transactions)
    raise FinancialValidationError('Choose csv, xlsx, or xls.')


def template_response(file_format):
    file_format = (file_format or '').lower()
    sample = {
        'transaction_id': '',
        'date': '2026-09-26',
        'transaction_type': 'EXPENSE',
        'amount': '100.00',
        'description': 'Example transaction',
        'status': 'ACTIVE',
        'person_id': '',
        'person_name': '',
        'account_id': '',
        'account_name': 'EXACT ACCOUNT NAME',
        'source_account_id': '',
        'source_account_name': '',
        'destination_account_id': '',
        'destination_account_name': '',
        'payment_source': '',
        'created_at': '',
    }
    if file_format == 'csv':
        response = HttpResponse(content_type='text/csv; charset=utf-8')
        response['Content-Disposition'] = 'attachment; filename="ledgerly_transaction_import_template.csv"'
        response.write('\ufeff')
        writer = csv.DictWriter(response, fieldnames=FORMAT_COLUMNS)
        writer.writeheader()
        writer.writerow(sample)
        return response
    if file_format == 'xlsx':
        return _template_xlsx(sample)
    if file_format == 'xls':
        return _template_xls(sample)
    raise FinancialValidationError('Choose csv, xlsx, or xls.')


def _template_xlsx(sample):
    if openpyxl is None:
        raise FinancialValidationError('XLSX support is unavailable. Install openpyxl.')
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = 'Transactions'
    sheet.append(list(FORMAT_COLUMNS))
    sheet.append([sample[column] for column in FORMAT_COLUMNS])
    sheet.freeze_panes = 'A2'
    for cell in sheet[1]:
        cell.font = openpyxl.styles.Font(bold=True)
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = min(max(len(str(column[0].value)) + 2, 14), 38)
    output = io.BytesIO()
    workbook.save(output)
    response = HttpResponse(output.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="ledgerly_transaction_import_template.xlsx"'
    return response


def _template_xls(sample):
    if xlwt is None:
        raise FinancialValidationError('XLS support is unavailable. Install xlwt.')
    workbook = xlwt.Workbook()
    sheet = workbook.add_sheet('Transactions')
    style = xlwt.easyxf('font: bold on; pattern: pattern solid, fore_colour gray25;')
    for col, name in enumerate(FORMAT_COLUMNS):
        sheet.write(0, col, name, style)
        sheet.write(1, col, sample[name])
        sheet.col(col).width = min(max(len(name) + 2, 14), 38) * 256
    output = io.BytesIO()
    workbook.save(output)
    response = HttpResponse(output.getvalue(), content_type='application/vnd.ms-excel')
    response['Content-Disposition'] = 'attachment; filename="ledgerly_transaction_import_template.xls"'
    return response
